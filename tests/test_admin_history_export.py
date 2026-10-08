from datetime import datetime

from app.config import normalize_telegram_chat_id
from app.export import (
    CHARGE_HEADERS,
    INSPECTION_HEADERS,
    charge_export_stats,
    custom_period_bounds,
    fp_link,
    period_bounds,
    write_inspections_xlsx,
)
from app.models import DamageControlCase, InspectionSession
from app.utils import is_supervisor, normalize_plate


def test_export_scores_supervisor_only():
    assert is_supervisor("Fedos_AV", "Fedos_AV")
    assert not is_supervisor("someone_else", "Fedos_AV")


def test_history_auto_normalizes_plate_before_lookup():
    assert normalize_plate("о917нх797") == "О917НХ797"


def test_period_bounds_today():
    start, end = period_bounds("today", datetime(2026, 6, 1, 12, 0))
    assert start == datetime(2026, 6, 1)
    assert end == datetime(2026, 6, 2)


def test_period_bounds_year():
    start, end = period_bounds("year", datetime(2026, 6, 3, 12, 0))
    assert start == datetime(2026, 1, 1)
    assert end == datetime(2027, 1, 1)


def test_period_bounds_quarter():
    start, end = period_bounds("quarter", datetime(2026, 6, 3, 12, 0))
    assert start == datetime(2026, 4, 1)
    assert end == datetime(2026, 7, 1)


def test_period_bounds_fourth_quarter_crosses_year():
    start, end = period_bounds("quarter", datetime(2026, 12, 3, 12, 0))
    assert start == datetime(2026, 10, 1)
    assert end == datetime(2027, 1, 1)


def test_period_bounds_all_time():
    start, end = period_bounds("all", datetime(2026, 6, 3, 12, 0))
    assert start == datetime(1970, 1, 1)
    assert end == datetime(2026, 6, 4)


def test_period_bounds_last_12_months():
    start, end = period_bounds("last12", datetime(2026, 7, 31, 12, 0))
    assert start == datetime(2025, 7, 31)
    assert end == datetime(2026, 8, 1)


def test_custom_period_uses_exclusive_next_day_boundary():
    start, end = custom_period_bounds("01.06.2026-30.06.2026")
    assert start == datetime(2026, 6, 1)
    assert end == datetime(2026, 7, 1)


def test_custom_period_rejects_reversed_dates():
    try:
        custom_period_bounds("30.06.2026-01.06.2026")
    except ValueError as exc:
        assert "Начальная дата" in str(exc)
    else:
        raise AssertionError(
            "Ожидалась ошибка для обратного периода"
        )


def test_charge_export_stats_separates_money_no_charge_and_legacy():
    stats = charge_export_stats(
        [
            DamageControlCase(payment_amount=5000),
            DamageControlCase(payment_amount=0),
            DamageControlCase(payment_amount=None),
        ]
    )
    assert stats == {"total": 3, "monetary": 1, "no_charge": 1, "legacy": 1, "amount": 5000}


def test_fp_link_from_positive_chat_id():
    row = InspectionSession(fp_chat_id=1001905865504, fp_message_id=123)
    assert fp_link(row) == "https://t.me/c/1905865504/123"


def test_positive_supergroup_chat_id_is_normalized_for_bot_api():
    assert normalize_telegram_chat_id("1001905865504") == -1001905865504
    assert normalize_telegram_chat_id("-1001905865504") == -1001905865504


def test_charge_export_headers_are_business_fields():
    assert CHARGE_HEADERS == [
        "ID списания",
        "Дата и время списания",
        "Номер авто",
        "ФИО водителя",
        "Тип повреждений",
        "Сумма списания",
        "Тип списания",
        "Комментарий",
    ]


def test_inspection_export_contains_id_and_all_business_fields(tmp_path):
    output = tmp_path / "inspections.xlsx"
    row = InspectionSession(
        id=42,
        plate_normalized="В981РН172",
        status="COMPLETED",
        completed_at=datetime(2026, 9, 25, 14, 30),
        scenario="PLANNED",
    )

    write_inspections_xlsx([row], output)

    from openpyxl import load_workbook

    ws = load_workbook(output).active
    assert [cell.value for cell in ws[1]] == INSPECTION_HEADERS
    assert ws[2][0].value == 42
    assert ws[2][1].value == "В981РН172"
